"""Build excel/ADMA_workbook.xlsx and its Ukrainian version excel/ADMA_workbook_ua.xlsx: every calculation of the
paper that can be checked by hand, written as live Excel formulas (not pasted values).

Sheets: README, Parameters, Calculator, Worked_example (start of Section 4), Statistics_mode (Tables 5 and 6),
TOPSIS (independent ranking of Table 6), EU_data, UA_data, Diagnostic_situations (Table 4),
Boundary_band (equation 7), Experiment_sample (first 1 000 profiles of the uniform scenario),
Figure_2_band and Figure_3_map (native Excel charts of Figures 2 and 3).

Inputs are read from data/raw/, data/matrix.csv, data/generated/ and results/simulation_experiment.json, so the workbook
always matches the scripts. excel/README.md (README_ua.md) is written from the same text as the README sheet.
The Ukrainian workbook has the same cells and formulas; only labels, sheet names and statements are translated
(texts in translations_ua.py).
Run scripts/simulation_experiment.py first.

Usage:  python scripts/build_workbook.py     (requires openpyxl)
"""

import csv
import io
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).resolve().parent))

from adma_index import CALIBRATED_PARAMETERS, ITEMS_INTEGRATION, ITEMS_PER_LAYER, ITEMS_STRATEGY, THEORETICAL_BOUNDS  # noqa: E402
from data_sources import (  # noqa: E402
    EU_AGGREGATE, EU_MEMBER_STATES, EUROSTAT_SDMX_URL, LAYERS, OUTSIDE_INDEX, RAW_EUROSTAT, REPOSITORY_ROOT, STATEMENT,
    STRENGTH, UKRSTAT_DATA_URL, INDICATOR_PAIRS, parse_value, read_eurostat, read_matrix)
from generate_profiles import GENERATED_DATA, column_names  # noqa: E402
from statistics_mode_local import COUNTRY_NAMES, RANKED, TABLE_ROWS  # noqa: E402
from translations_ua import (  # noqa: E402
    COUNTRY_NAMES_UA, README_INTRO_UA, README_SECTIONS_UA, README_TITLE_UA, SHEET_NAMES_UA, STATEMENTS_UA, STRENGTH_UA,
    WORKBOOK_TEXT_UA, translate_note)

WORKBOOKS = {  # language: (workbook, README written from the README sheet)
    "en": (REPOSITORY_ROOT / "excel" / "ADMA_workbook.xlsx", REPOSITORY_ROOT / "excel" / "README.md"),
    "ua": (REPOSITORY_ROOT / "excel" / "ADMA_workbook_ua.xlsx", REPOSITORY_ROOT / "excel" / "README_ua.md"),
}
SIMULATION_RESULTS = REPOSITORY_ROOT / "results" / "simulation_experiment.json"
SAMPLE_SIZE = 1000

BOLD = Font(bold=True)
TITLE = Font(bold=True, size=13)
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")  # yellow: values a reviewer may change
SOURCE_FILL = PatternFill("solid", fgColor="DDEBF7")  # blue: values copied from official sources
WRAP = Alignment(wrap_text=True, vertical="top")

STATEMENTS = [
    ("C", "Software for production and resource planning (ERP, MES)"),
    ("C", "Automation of production processes with digital equipment"),
    ("C", "Approved digital development strategy or plan for 2-3 years"),
    ("L1", "Digital tools for human resource management (HRM)"),
    ("L1", "Employees have the skills to work with digital tools"),
    ("L1", "Production staff enter primary data on mobile devices or terminals at the workplace"),
    ("L2", "Digital accounting of inventories (WMS or spreadsheets)"),
    ("L2", "Collection and analysis of data on resource use, deviations and losses"),
    ("L2", "Monitoring of storage conditions with sensors or IoT"),
    ("L3", "GPS monitoring or telematics on vehicles and machinery"),
    ("L3", "Data-based maintenance planning (CMMS)"),
    ("L3", "Precision farming or precision livestock technologies"),
    ("L4", "Digital sales channels (website, e-commerce, EDI)"),
    ("L4", "Systematic collection and analysis of demand, price and feedback data (CRM)"),
    ("L4", "Digital traceability: verified information on product origin"),
    ("I", "Automatic transfer of production data between units"),
    ("I", "Timely receipt of external data (prices, weather, regulation)"),
    ("I", "Integration of the enterprise's digital systems with each other"),
    ("S", "Regular tracking of new digital technologies"),
    ("S", "Dedicated budget or resources for digitalisation"),
    ("S", "Evaluation of the results of digital projects and revision of digitalisation plans"),
]
# Worked example at the start of Section 4: C = 7/12, L1 = L2 = L4 = 0.5, L3 = 1/12.
WORKED_EXAMPLE_ANSWERS = [3, 3, 4, 3, 3, 3, 3, 3, 3, 2, 1, 1, 3, 3, 3, 2, 2, 2, 2, 2, 2]
STATISTICS_INDICATORS = [code for codes in list(LAYERS.values()) + list(OUTSIDE_INDEX.values()) for code in codes]
INDICATOR_LAYER = {code: layer for layer, codes in list(LAYERS.items()) + list(OUTSIDE_INDEX.items()) for code in codes}
PAIRS = {pair["pair_id"]: pair for pair in INDICATOR_PAIRS}


def add_name(workbook, name, reference):
    workbook.defined_names[name] = DefinedName(name, attr_text=reference)


def level_formula(cell):
    """Level 1..5; a value exactly on a bound belongs to the lower level (same rule as the scripts)."""
    return '=1+COUNTIF(bounds,"<"&ROUND(%s,12))' % cell


def set_widths(sheet, widths):
    for column_letter, width in widths.items():
        sheet.column_dimensions[column_letter].width = width


language = "en"  # language of the workbook being built, set by set_language()


def set_language(code):
    global language
    language = code


def tr(text):
    """Label in the language of the workbook; every English label has an entry in WORKBOOK_TEXT_UA."""
    return text if language == "en" else WORKBOOK_TEXT_UA[text]


def sheet_name(name):
    return name if language == "en" else SHEET_NAMES_UA[name]


def ref(name):
    """Prefix of a reference to another sheet in a formula (quoted, since Ukrainian names are not plain ASCII)."""
    return "%s!" % name if language == "en" else "'%s'!" % SHEET_NAMES_UA[name]


def country_name(country):
    return (COUNTRY_NAMES if language == "en" else COUNTRY_NAMES_UA)[country]


def statements():
    return STATEMENTS if language == "en" else [(block, text) for (block, _), text in zip(STATEMENTS, STATEMENTS_UA)]


README_TITLE = "ADMA workbook: hand-checkable calculations of the paper"
README_INTRO = [
    "Every number on the calculation sheets is an Excel formula. Yellow cells are inputs you may change; "
    "blue cells are values copied from official sources (each with its link). Changing the yellow cells on "
    "Parameters recalculates every sheet.",
    "A Ukrainian version with the same cells and formulas is excel/ADMA_workbook_ua.xlsx.",
]
README_SECTIONS = [
    ("Sheets and the paper", [
        "Parameters: alpha, delta and level bounds used by all formulas (Section 3, Table 2).",
        "Calculator: enter the 21 questionnaire answers (1-5) and get layer scores, penalty, ADMI, level, DM_I, DM_S (equations 1-5).",
        "Worked_example: the profile at the start of Section 4.",
        "Statistics_mode: Tables 5 and 6, food processing (NACE C10-C12) of EU-27, Germany, nine CEE countries and Ukraine (equation 6, P_min = 0). Uses EU_data and UA_data.",
        "TOPSIS: the independent ranking of Table 6 (entropy weights, closeness to the ideal solution) and its Spearman correlation with ADMI.",
        "EU_data: Eurostat values for the EU-27 aggregate and all 27 member states; the leader is found with MAX/INDEX/MATCH.",
        "UA_data: Derzhstat values with direct download links.",
        "Diagnostic_situations: Table 4.",
        "Boundary_band: equation (7), the band of base index values in which the penalty changes the level (P1).",
        "Experiment_sample: the first 1 000 of the 100 000 generated profiles (uniform scenario) with all formulas row by row, "
        "and the corresponding results of the full run for comparison.",
        "Figure_2_band: native Excel chart of the band of equation (7) with the sample profiles whose level the penalty changed.",
        "Figure_3_map: native Excel chart of the diagnostic map (ADMI, DM_I) with the situations of Table 4.",
    ]),
    ("Full experiment and statistics mode", [
        "The complete experiment (2 x 100 000 profiles, 99 parameter combinations, Table 3, Supplementary Table S2, propositions P1-P3) is run by scripts/simulation_experiment.py.",
        "The statistics mode with the proxy robustness check (Tables 5 and 6, Supplementary Table S3, proposition P4) is run by scripts/statistics_mode_local.py.",
        "The profiles are stored in data/generated/profiles_uniform.csv and data/generated/profiles_low_skewed.csv; "
        "any row can be typed into the Calculator to check it.",
    ]),
    ("Level rule", [
        "Level = 1 + number of bounds strictly below the index value rounded to 12 decimals. A value exactly on a bound "
        "belongs to the lower level: 0.40 is level 2, 0.41 is level 3.",
        "Index values computed from questionnaire answers are multiples of 1/60, so exact ties occur; the rounding removes "
        "floating-point noise so that a tie is never moved to the upper level.",
    ]),
    ("Verification", [
        "scripts/verify_workbook.py evaluates the formulas of this workbook with an independent engine (pycel) and compares "
        "7 533 values with the Python scripts. All values agree.",
        "The workbook is generated by scripts/build_workbook.py from data/raw, data/matrix.csv and data/generated; this README is generated "
        "together with the README sheet and has the same content.",
    ]),
]


def readme_text():
    return (README_TITLE, README_INTRO, README_SECTIONS) if language == "en" else (README_TITLE_UA, README_INTRO_UA, README_SECTIONS_UA)


def build_readme(sheet):
    title, intro, sections = readme_text()
    sheet["A1"] = title
    sheet["A1"].font = TITLE
    row_number = 3
    for paragraph in intro:
        sheet.cell(row=row_number, column=1, value=paragraph)
        row_number += 1
    for heading, lines in sections:
        row_number += 1
        sheet.cell(row=row_number, column=1, value=heading).font = BOLD
        row_number += 1
        for line in lines:
            sheet.cell(row=row_number, column=1, value=line)
            row_number += 1
    set_widths(sheet, {"A": 160})


def write_readme_markdown(path):
    """Same content as the README sheet, for reading on GitHub without opening Excel."""
    title, intro, sections = readme_text()
    lines = ["# " + title, ""]
    lines.extend(intro)
    for heading, section_lines in sections:
        lines.extend(["", "## " + heading, ""])
        lines.extend("- " + line for line in section_lines)
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def build_parameters(workbook):
    sheet = workbook.create_sheet(sheet_name("Parameters"))
    sheet["A1"] = tr("Parameters")
    sheet["A1"].font = TITLE
    rows = [
        ("alpha", CALIBRATED_PARAMETERS["alpha"], tr("Strength of the weak-link penalty (calibrated value)")),
        ("delta", CALIBRATED_PARAMETERS["delta"], tr("Threshold of critical layer weakness (calibrated value; 0.375 = mean answer 2.5)")),
    ]
    for row_number, (name, value, description) in enumerate(rows, start=3):
        sheet.cell(row=row_number, column=1, value=name).font = BOLD
        sheet.cell(row=row_number, column=2, value=value).fill = INPUT_FILL
        sheet.cell(row=row_number, column=3, value=description)
    add_name(workbook, "alpha", ref("Parameters") + "$B$3")
    add_name(workbook, "delta", ref("Parameters") + "$B$4")
    sheet["A6"] = tr("Level bounds (Table 2)")
    sheet["A6"].font = BOLD
    labels = [tr("Level 1 | 2"), tr("Level 2 | 3"), tr("Level 3 | 4"), tr("Level 4 | 5")]
    for offset, (label, bound) in enumerate(zip(labels, THEORETICAL_BOUNDS)):
        sheet.cell(row=7 + offset, column=1, value=label)
        sheet.cell(row=7 + offset, column=2, value=bound).fill = INPUT_FILL
    add_name(workbook, "bounds", ref("Parameters") + "$B$7:$B$10")
    sheet["A12"] = tr("Initial values before calibration: alpha = 0.50, delta = 0.25. Model-adjusted bounds used as a robustness check: 0.25, 0.45, 0.60, 0.80.")
    set_widths(sheet, {"A": 22, "B": 10, "C": 80})


def build_calculator(workbook, name, answers, title):
    sheet = workbook.create_sheet(sheet_name(name))
    sheet["A1"] = tr(title)
    sheet["A1"].font = TITLE
    sheet.append([])
    sheet.append(["#", tr("Block"), tr("Statement"), tr("Answer (1-5)"), tr("Normalised (eq. 1)")])
    for cell in sheet[3]:
        cell.font = BOLD
    validation = DataValidation(type="whole", operator="between", formula1="1", formula2="5", showErrorMessage=True)
    sheet.add_data_validation(validation)
    first_row = 4
    for index, ((block, text), answer) in enumerate(zip(statements(), answers)):
        row = first_row + index
        sheet.cell(row=row, column=1, value=index + 1)
        sheet.cell(row=row, column=2, value=block)
        sheet.cell(row=row, column=3, value=text)
        answer_cell = sheet.cell(row=row, column=4, value=answer)
        answer_cell.fill = INPUT_FILL
        validation.add(answer_cell)
        sheet.cell(row=row, column=5, value="=(D%d-1)/4" % row)
    last_row = first_row + len(STATEMENTS) - 1
    summary_row = last_row + 2
    sheet.cell(row=summary_row, column=3, value=tr("Result")).font = BOLD
    block_rows = {}
    for index, (block, _text) in enumerate(STATEMENTS):
        block_rows.setdefault(block, []).append(first_row + index)
    results = []
    for block in list(ITEMS_PER_LAYER) + ["I", "S"]:
        rows = block_rows[block]
        label = {"I": tr("DM_I (integration, outside the index)"), "S": tr("DM_S (strategy, outside the index)")}.get(block, tr("DM_%s (eq. 2)") % block)
        results.append((label, "=AVERAGE(E%d:E%d)" % (rows[0], rows[-1])))
    layer_range = "D%d:D%d" % (summary_row + 1, summary_row + len(ITEMS_PER_LAYER))
    base_row = summary_row + 1 + len(results)
    results.extend([
        (tr("ADMI_base (eq. 3)"), "=AVERAGE(%s)" % layer_range),
        ("DM_min", "=MIN(%s)" % layer_range),
        (tr("Penalty P (eq. 4)"), "=1-alpha*MAX(0,delta-D%d)" % (base_row + 1)),
        (tr("ADMI (eq. 5)"), "=D%d*D%d" % (base_row, base_row + 2)),
        (tr("Level by ADMI_base"), level_formula("D%d" % base_row)),
        (tr("Level by ADMI"), level_formula("D%d" % (base_row + 3))),
        (tr("Penalty changed the level?"), '=IF(D%d<>D%d,"%s","%s")' % (base_row + 4, base_row + 5, tr("yes"), tr("no"))),
    ])
    for offset, (label, formula) in enumerate(results, start=1):
        sheet.cell(row=summary_row + offset, column=3, value=label)
        sheet.cell(row=summary_row + offset, column=4, value=formula)
    set_widths(sheet, {"A": 4, "B": 7, "C": 78, "D": 14, "E": 18})
    return sheet


def column_letter(column):
    from openpyxl.utils import get_column_letter
    return get_column_letter(column)


EU_FIRST_ROW = 7  # EU-27 aggregate on row 6, member states on rows 7..33
EU_LEADER_VALUE_ROW = EU_FIRST_ROW + len(EU_MEMBER_STATES) + 1
EU_LEADER_ROW = EU_LEADER_VALUE_ROW + 1
UA_FIRST_ROW = 4


def build_data_sheets(workbook):
    """EU_data (EU-27 aggregate and all member states) and UA_data, one column or row per indicator."""
    matrix = {row["pair_id"]: row for row in read_matrix()}
    eu_sheet = workbook.create_sheet(sheet_name("EU_data"))
    eu_sheet["A1"] = tr("Eurostat, ICT usage in enterprises: NACE C10-C12, enterprises with 10+ persons employed, % of enterprises")
    eu_sheet["A1"].font = TITLE
    for row, label in ((2, tr("Indicator")), (3, tr("Statement")), (4, tr("Year")), (5, tr("Source (CSV)")), (6, EU_AGGREGATE),
                       (EU_LEADER_VALUE_ROW, tr("Leader value")), (EU_LEADER_ROW, tr("Leader"))):
        eu_sheet.cell(row=row, column=1, value=label).font = BOLD
    for offset, country in enumerate(EU_MEMBER_STATES):
        eu_sheet.cell(row=EU_FIRST_ROW + offset, column=1, value=country)
    last_member_row = EU_FIRST_ROW + len(EU_MEMBER_STATES) - 1
    eurostat_cache = {}
    for offset, code in enumerate(STATISTICS_INDICATORS):
        column = 2 + offset
        letter = column_letter(column)
        pair = PAIRS[code]
        year = matrix[code]["eurostat_year_used"]
        if pair["eurostat_dataset"] not in eurostat_cache:
            eurostat_cache[pair["eurostat_dataset"]] = read_eurostat(RAW_EUROSTAT, pair["eurostat_dataset"])
        observations = eurostat_cache[pair["eurostat_dataset"]]
        eu_sheet.cell(row=2, column=column, value="%s (%s)" % (code, pair["eurostat_indicator"])).font = BOLD
        eu_sheet.cell(row=3, column=column, value=STATEMENT[code])
        eu_sheet.cell(row=4, column=column, value=int(year))
        eu_sheet.cell(row=5, column=column, value=EUROSTAT_SDMX_URL.format(dataset=pair["eurostat_dataset"], indicator=pair["eurostat_indicator"]))
        for row, country in [(6, EU_AGGREGATE)] + [(EU_FIRST_ROW + index, member) for index, member in enumerate(EU_MEMBER_STATES)]:
            observation = observations.get((pair["eurostat_indicator"], country, year))
            value = parse_value(observation["value"]) if observation else None
            eu_sheet.cell(row=row, column=column, value=value).fill = SOURCE_FILL
        member_range = "%s%d:%s%d" % (letter, EU_FIRST_ROW, letter, last_member_row)
        eu_sheet.cell(row=EU_LEADER_VALUE_ROW, column=column, value="=MAX(%s)" % member_range)
        eu_sheet.cell(row=EU_LEADER_ROW, column=column, value="=INDEX($A$%d:$A$%d,MATCH(MAX(%s),%s,0))" % (
            EU_FIRST_ROW, last_member_row, member_range, member_range))
    eu_sheet.cell(row=EU_LEADER_ROW + 2, column=1, value=tr("Empty cell = no published value (for example flag u, low reliability). "
                  "Leader = member state with the highest published value. The CSV links download the full series; filter by geo and TIME_PERIOD."))
    set_widths(eu_sheet, {"A": 14})

    ua_sheet = workbook.create_sheet(sheet_name("UA_data"))
    ua_sheet["A1"] = tr("State Statistics Service of Ukraine, data bank stat.gov.ua (SSSU:DF_INFORM_COMMUN_TECH_ENTRP): KVED 10-12, enterprises 10+, % of enterprises")
    ua_sheet["A1"].font = TITLE
    ua_sheet.append([])
    ua_sheet.append([tr(header) for header in ["Indicator", "Statement", "Derzhstat series", "Year", "Value, %", "Note", "Direct download (CSV)"]])
    for cell in ua_sheet[3]:
        cell.font = BOLD
    for offset, code in enumerate(STATISTICS_INDICATORS):
        row = UA_FIRST_ROW + offset
        matrix_row = matrix[code]
        ua_sheet.cell(row=row, column=1, value=code)
        ua_sheet.cell(row=row, column=2, value=STATEMENT[code])
        ua_sheet.cell(row=row, column=3, value=matrix_row["ukrstat_series"])
        ua_sheet.cell(row=row, column=4, value=int(matrix_row["ukrstat_year_used"]))
        ua_sheet.cell(row=row, column=5, value=parse_value(matrix_row["UA_value"])).fill = SOURCE_FILL
        note = matrix_row["UA_flag_or_note"]
        ua_sheet.cell(row=row, column=6, value=note if language == "en" else translate_note(note))
        ua_sheet.cell(row=row, column=7, value=UKRSTAT_DATA_URL.format(indicator=PAIRS[code]["ukrstat_indicator"]))
    ua_sheet.cell(row=UA_FIRST_ROW + len(STATISTICS_INDICATORS) + 1, column=1,
                  value=tr("Each link downloads a CSV with all years and breakdowns (columns BREAKDOWN_CATEGORY, BREAKDOWN, TIME_PERIOD, OBS_VALUE). "
                           "Interactive table: https://stat.gov.ua/uk/explorer?urn=SSSU%3ADF_INFORM_COMMUN_TECH_ENTRP%28~%29"))
    set_widths(ua_sheet, {"A": 30, "B": 10, "C": 46, "D": 7, "E": 10, "F": 60, "G": 120})


STATISTICS_FIRST_ROW = 4
STATISTICS_LAST_ROW = STATISTICS_FIRST_ROW + len(STATISTICS_INDICATORS) - 1
OBJECT_FIRST_COLUMN = 9  # column I
INDEX_HEADER_ROW = STATISTICS_LAST_ROW + 3
INDEX_FIRST_ROW = INDEX_HEADER_ROW + 1
INDEX_COLUMNS = ["Object", "C", "L1", "L2", "L3", "L4", "ADMI_base", "DM_min", "Penalty", "ADMI", "Level", "DM_I", "DM_S",
                 "Rank ADMI", "Rank TOPSIS"]


def index_row(country):
    return INDEX_FIRST_ROW + TABLE_ROWS.index(country)


def index_cell(country, header):
    return "%s%d" % (column_letter(INDEX_COLUMNS.index(header) + 1), index_row(country))


def build_statistics_mode(workbook):
    sheet = workbook.create_sheet(sheet_name("Statistics_mode"))
    sheet["A1"] = tr("Tables 5 and 6. Statistics mode of ADMA for food processing (NACE C10-C12): layer scores, index and ranking (equation 6, P_min = 0)")
    sheet["A1"].font = TITLE
    sheet["A2"] = tr("Upper part: every indicator divided by the value of the EU leader. Lower part: layers, ADMI and ranks. EU-27 and Germany are reference rows outside the ranking.")
    headers = [tr(header) for header in ["Layer", "Indicator", "Statement", "Strength", "EU year", "UA year", "Leader", "Leader, %"]] + [
        country_name(country) for country in TABLE_ROWS]
    for column, header in enumerate(headers, start=1):
        sheet.cell(row=3, column=column, value=header).font = BOLD
    member_rows = {member: EU_FIRST_ROW + index for index, member in enumerate(EU_MEMBER_STATES)}
    member_rows[EU_AGGREGATE] = 6
    for offset, code in enumerate(STATISTICS_INDICATORS):
        row = STATISTICS_FIRST_ROW + offset
        eu_letter = column_letter(2 + offset)
        sheet.cell(row=row, column=1, value=INDICATOR_LAYER[code])
        sheet.cell(row=row, column=2, value=code)
        sheet.cell(row=row, column=3, value=STATEMENT[code])
        sheet.cell(row=row, column=4, value=STRENGTH[code] if language == "en" else STRENGTH_UA[STRENGTH[code]])
        sheet.cell(row=row, column=5, value="=%s%s4" % (ref("EU_data"), eu_letter))
        sheet.cell(row=row, column=6, value="=%sD%d" % (ref("UA_data"), UA_FIRST_ROW + offset))
        sheet.cell(row=row, column=7, value="=%s%s%d" % (ref("EU_data"), eu_letter, EU_LEADER_ROW))
        sheet.cell(row=row, column=8, value="=%s%s%d" % (ref("EU_data"), eu_letter, EU_LEADER_VALUE_ROW))
        for country_offset, country in enumerate(TABLE_ROWS):
            source = "%sE%d" % (ref("UA_data"), UA_FIRST_ROW + offset) if country == "UA" else "%s%s%d" % (ref("EU_data"), eu_letter, member_rows[country])
            sheet.cell(row=row, column=OBJECT_FIRST_COLUMN + country_offset, value='=IF(%s="","",%s/$H%d)' % (source, source, row))
    for column, header in enumerate(INDEX_COLUMNS, start=1):
        label = tr(header) if header in ("Object", "Penalty", "Level", "Rank ADMI", "Rank TOPSIS") else header
        sheet.cell(row=INDEX_HEADER_ROW, column=column, value=label).font = BOLD
    ranked_first, ranked_last = index_row(RANKED[0]), index_row(RANKED[-1])
    for country_offset, country in enumerate(TABLE_ROWS):
        row = index_row(country)
        object_letter = column_letter(OBJECT_FIRST_COLUMN + country_offset)
        scores = "%s$%d:%s$%d" % (object_letter, STATISTICS_FIRST_ROW, object_letter, STATISTICS_LAST_ROW)
        sheet.cell(row=row, column=1, value=country_name(country))
        for header in ["C", "L1", "L2", "L3", "L4", "DM_I", "DM_S"]:
            layer = header[3:] if header.startswith("DM_") else header
            sheet.cell(row=row, column=INDEX_COLUMNS.index(header) + 1,
                       value='=AVERAGEIF($A$%d:$A$%d,"%s",%s)' % (STATISTICS_FIRST_ROW, STATISTICS_LAST_ROW, layer, scores))
        layer_range = "B%d:F%d" % (row, row)
        sheet.cell(row=row, column=7, value="=AVERAGE(%s)" % layer_range)
        sheet.cell(row=row, column=8, value="=MIN(%s)" % layer_range)
        sheet.cell(row=row, column=9, value="=1-alpha*MAX(0,delta-H%d)" % row)
        sheet.cell(row=row, column=10, value="=G%d*I%d" % (row, row))
        sheet.cell(row=row, column=11, value=level_formula("J%d" % row))
        if country in RANKED:
            sheet.cell(row=row, column=14, value='=1+COUNTIF($J$%d:$J$%d,">"&J%d)' % (ranked_first, ranked_last, row))
            sheet.cell(row=row, column=15, value="=%s%s%d" % (ref("TOPSIS"), TOPSIS_COLUMNS["rank"], TOPSIS_FIRST_ROW + RANKED.index(country)))
        else:
            sheet.cell(row=row, column=14, value="-")
            sheet.cell(row=row, column=15, value="-")
    notes = [
        "Strength of correspondence to the questionnaire statement: direct, acceptable or proxy (Supplementary Table S1).",
        "Years differ for some pairs (Supplementary Table S1). EU-27 and Germany do not enter the ranking.",
        "The strict set without weak proxies and the leave-one-out check (Supplementary Table S3) are computed by scripts/statistics_mode_local.py.",
    ]
    for offset, note in enumerate(notes):
        sheet.cell(row=INDEX_FIRST_ROW + len(TABLE_ROWS) + 1 + offset, column=1, value=tr(note))
    set_widths(sheet, {"A": 12, "B": 30, "C": 9, "D": 11, "E": 8, "F": 8, "G": 8, "H": 9})


TOPSIS_FIRST_ROW = 5
TOPSIS_LAYERS = list(LAYERS)
TOPSIS_BLOCKS = {"score": 2, "share": 8, "entropy_term": 14, "weighted": 20}  # first column of each five-column block
TOPSIS_COLUMNS = {"to_ideal": "Z", "to_anti_ideal": "AA", "closeness": "AB", "rank": "AC", "admi_rank": "AD", "squared_difference": "AE"}


def build_topsis(workbook):
    """TOPSIS with entropy weights on the five layers of Table 5 (Hwang and Yoon, 1981), every step a formula."""
    sheet = workbook.create_sheet(sheet_name("TOPSIS"))
    sheet["A1"] = tr("TOPSIS ranking of the ten ranked objects of Table 6 on the same five layers, entropy weights")
    sheet["A1"].font = TITLE
    last_row = TOPSIS_FIRST_ROW + len(RANKED) - 1
    count = len(RANKED)
    titles = {"score": "Layer score", "share": "Share p = x / column sum", "entropy_term": "p * ln p", "weighted": "Weighted normalised v = w * x / sqrt(sum x^2)"}
    for block, first_column in TOPSIS_BLOCKS.items():
        sheet.cell(row=3, column=first_column, value=tr(titles[block])).font = BOLD
        for layer_offset, layer in enumerate(TOPSIS_LAYERS):
            sheet.cell(row=4, column=first_column + layer_offset, value=layer).font = BOLD
    sheet.cell(row=4, column=1, value=tr("Object")).font = BOLD
    for header, letter in (("S+ (to ideal)", "Z"), ("S- (to anti-ideal)", "AA"), ("Closeness", "AB"), ("Rank TOPSIS", "AC"),
                           ("Rank ADMI", "AD"), ("(rank difference)^2", "AE")):
        sheet["%s4" % letter] = tr(header)
        sheet["%s4" % letter].font = BOLD
    summary_rows = {"entropy": last_row + 2, "divergence": last_row + 3, "weight": last_row + 4, "ideal": last_row + 5, "anti_ideal": last_row + 6}
    for name, row in summary_rows.items():
        sheet.cell(row=row, column=1, value=tr({"entropy": "Entropy e", "divergence": "Divergence 1 - e", "weight": "Entropy weight w",
                                                "ideal": "Ideal (max v)", "anti_ideal": "Anti-ideal (min v)"}[name])).font = BOLD
    for offset, country in enumerate(RANKED):
        row = TOPSIS_FIRST_ROW + offset
        sheet.cell(row=row, column=1, value=country_name(country))
        for layer_offset, layer in enumerate(TOPSIS_LAYERS):
            score = column_letter(TOPSIS_BLOCKS["score"] + layer_offset)
            share = column_letter(TOPSIS_BLOCKS["share"] + layer_offset)
            weighted_column = TOPSIS_BLOCKS["weighted"] + layer_offset
            sheet.cell(row=row, column=TOPSIS_BLOCKS["score"] + layer_offset, value="=%s%s" % (ref("Statistics_mode"), index_cell(country, layer)))
            sheet.cell(row=row, column=TOPSIS_BLOCKS["share"] + layer_offset, value="=%s%d/SUM(%s$%d:%s$%d)" % (score, row, score, TOPSIS_FIRST_ROW, score, last_row))
            sheet.cell(row=row, column=TOPSIS_BLOCKS["entropy_term"] + layer_offset, value="=IF(%s%d>0,%s%d*LN(%s%d),0)" % (share, row, share, row, share, row))
            sheet.cell(row=row, column=weighted_column, value="=%s$%d*%s%d/SQRT(SUMPRODUCT(%s$%d:%s$%d,%s$%d:%s$%d))" % (
                score, summary_rows["weight"], score, row, score, TOPSIS_FIRST_ROW, score, last_row, score, TOPSIS_FIRST_ROW, score, last_row))
        weighted = [column_letter(TOPSIS_BLOCKS["weighted"] + layer_offset) for layer_offset in range(len(TOPSIS_LAYERS))]
        sheet["Z%d" % row] = "=SQRT(%s)" % "+".join("(%s%d-%s$%d)^2" % (letter, row, letter, summary_rows["ideal"]) for letter in weighted)
        sheet["AA%d" % row] = "=SQRT(%s)" % "+".join("(%s%d-%s$%d)^2" % (letter, row, letter, summary_rows["anti_ideal"]) for letter in weighted)
        sheet["AB%d" % row] = "=AA%d/(Z%d+AA%d)" % (row, row, row)
        sheet["AC%d" % row] = '=1+COUNTIF($AB$%d:$AB$%d,">"&AB%d)' % (TOPSIS_FIRST_ROW, last_row, row)
        sheet["AD%d" % row] = "=%s%s" % (ref("Statistics_mode"), index_cell(country, "Rank ADMI"))
        sheet["AE%d" % row] = "=(AC%d-AD%d)^2" % (row, row)
    for layer_offset in range(len(TOPSIS_LAYERS)):
        score = column_letter(TOPSIS_BLOCKS["score"] + layer_offset)
        term = column_letter(TOPSIS_BLOCKS["entropy_term"] + layer_offset)
        weighted = column_letter(TOPSIS_BLOCKS["weighted"] + layer_offset)
        first_score, last_score = column_letter(TOPSIS_BLOCKS["score"]), column_letter(TOPSIS_BLOCKS["score"] + len(TOPSIS_LAYERS) - 1)
        sheet["%s%d" % (score, summary_rows["entropy"])] = "=-SUM(%s%d:%s%d)/LN(%d)" % (term, TOPSIS_FIRST_ROW, term, last_row, count)
        sheet["%s%d" % (score, summary_rows["divergence"])] = "=1-%s%d" % (score, summary_rows["entropy"])
        sheet["%s%d" % (score, summary_rows["weight"])] = "=%s%d/SUM($%s$%d:$%s$%d)" % (
            score, summary_rows["divergence"], first_score, summary_rows["divergence"], last_score, summary_rows["divergence"])
        sheet["%s%d" % (weighted, summary_rows["ideal"])] = "=MAX(%s%d:%s%d)" % (weighted, TOPSIS_FIRST_ROW, weighted, last_row)
        sheet["%s%d" % (weighted, summary_rows["anti_ideal"])] = "=MIN(%s%d:%s%d)" % (weighted, TOPSIS_FIRST_ROW, weighted, last_row)
    spearman_row = last_row + 8
    sheet.cell(row=spearman_row, column=1, value=tr("Spearman rank correlation ADMI - TOPSIS")).font = BOLD
    sheet.cell(row=spearman_row, column=2, value="=1-6*SUM(AE%d:AE%d)/(%d*(%d^2-1))" % (TOPSIS_FIRST_ROW, last_row, count, count))
    set_widths(sheet, {"A": 34})


def build_diagnostic_situations(workbook):
    sheet = workbook.create_sheet(sheet_name("Diagnostic_situations"))
    sheet["A1"] = tr("Table 4. Diagnostic situations of ADMA (illustrative profiles)")
    sheet["A1"].font = TITLE
    sheet.append([])
    sheet.append([tr("Situation"), "DM_C", "DM_L1", "DM_L2", "DM_L3", "DM_L4", "DM_I", "DM_S", "ADMI_base", "DM_min", tr("Penalty"), "ADMI",
                  tr("Level by ADMI"), tr("Scalar incl. DM_I"), tr("Level by scalar")])
    for cell in sheet[3]:
        cell.font = BOLD
    situations = [
        ("Compliance trap", [0.70, 0.40, 0.45, 0.05, 0.60], 0.30, 0.40),
        ("External EU pressure", [0.35, 0.30, 0.35, 0.30, 0.70], 0.15, 0.30),
        ("Absorptive limiter", [0.55, 0.10, 0.45, 0.45, 0.50], 0.35, 0.40),
        ("Aware mismatch", [0.30, 0.25, 0.25, 0.30, 0.25], 0.20, 0.75),
        ("Illusion of maturity", [0.55, 0.55, 0.55, 0.55, 0.55], 0.15, 0.50),
    ]
    for offset, (name, layers, integration, strategy) in enumerate(situations):
        row = 4 + offset
        sheet.cell(row=row, column=1, value=tr(name))
        for layer_offset, value in enumerate(layers + [integration, strategy]):
            sheet.cell(row=row, column=2 + layer_offset, value=value).fill = INPUT_FILL
        sheet.cell(row=row, column=9, value="=AVERAGE(B%d:F%d)" % (row, row))
        sheet.cell(row=row, column=10, value="=MIN(B%d:F%d)" % (row, row))
        sheet.cell(row=row, column=11, value="=1-alpha*MAX(0,delta-J%d)" % row)
        sheet.cell(row=row, column=12, value="=I%d*K%d" % (row, row))
        sheet.cell(row=row, column=13, value=level_formula("L%d" % row))
        sheet.cell(row=row, column=14, value="=AVERAGE(B%d:G%d)" % (row, row))
        sheet.cell(row=row, column=15, value=level_formula("N%d" % row))
    set_widths(sheet, {"A": 22})


def build_boundary_band(workbook):
    sheet = workbook.create_sheet(sheet_name("Boundary_band"))
    sheet["A1"] = tr("Equation (7): the level drops across bound b if and only if b < ADMI_base <= b / P(DM_min)")
    sheet["A1"].font = TITLE
    sheet["A2"] = tr("Upper limit is also capped by (4 + DM_min) / 5, the largest base index possible when the weakest layer equals DM_min.")
    sheet.append([tr("Bound b"), "DM_min", tr("Penalty P"), tr("ADMI_base from"), tr("ADMI_base to"), tr("Band width")])
    for cell in sheet[3]:
        cell.font = BOLD
    row = 4
    for bound_offset in range(len(THEORETICAL_BOUNDS)):
        for step in range(16):
            sheet.cell(row=row, column=1, value="=INDEX(bounds,%d)" % (bound_offset + 1))
            sheet.cell(row=row, column=2, value=round(step * 0.025, 3))
            sheet.cell(row=row, column=3, value="=1-alpha*MAX(0,delta-B%d)" % row)
            sheet.cell(row=row, column=4, value="=A%d" % row)
            sheet.cell(row=row, column=5, value="=MIN(A%d/C%d,(4+B%d)/5)" % (row, row, row))
            sheet.cell(row=row, column=6, value="=MAX(0,E%d-A%d)" % (row, row))
            row += 1


def build_experiment_sample(workbook):
    sheet = workbook.create_sheet(sheet_name("Experiment_sample"))
    names = column_names()
    path = GENERATED_DATA / "profiles_uniform.csv"
    with io.open(str(path), encoding="utf-8") as handle:
        reader = csv.reader(handle)
        next(reader)
        sample = [next(reader) for _ in range(SAMPLE_SIZE)]
    sheet["A1"] = tr("First %d of 100 000 generated profiles (uniform scenario, data/generated/profiles_uniform.csv), every column is a formula") % SAMPLE_SIZE
    sheet["A1"].font = TITLE
    headers = ["profile_id"] + names + ["DM_C", "DM_L1", "DM_L2", "DM_L3", "DM_L4", "DM_I", "DM_S", "ADMI_base", "DM_min", "Penalty", "ADMI",
                                        "Level base", "Level ADMI", "Penalised", "Level changed", "Illusion of maturity", "Scalar incl. DM_I", "Scalar level",
                                        "Figure 2: ADMI_base if level changed", "Figure 2: ADMI_base if level kept"]
    header_row = 8
    for column, header in enumerate(headers, start=1):
        sheet.cell(row=header_row, column=column, value=sample_label(header)).font = BOLD
    first_row = header_row + 1
    last_row = first_row + SAMPLE_SIZE - 1
    answer_columns = {name: sheet.cell(row=1, column=2 + index).column_letter for index, name in enumerate(names)}

    def block_range(prefix, item_count, row):
        return "%s%d:%s%d" % (answer_columns["%s_1" % prefix], row, answer_columns["%s_%d" % (prefix, item_count)], row)

    blocks = [(layer, count) for layer, count in ITEMS_PER_LAYER.items()] + [("I", ITEMS_INTEGRATION), ("S", ITEMS_STRATEGY)]
    result_start = 2 + len(names)
    letters = {header: sheet.cell(row=1, column=column).column_letter for column, header in enumerate(headers, start=1)}
    for offset, values in enumerate(sample):
        row = first_row + offset
        for column, value in enumerate(values, start=1):
            sheet.cell(row=row, column=column, value=int(value))
        for block_offset, (prefix, count) in enumerate(blocks):
            sheet.cell(row=row, column=result_start + block_offset, value="=(AVERAGE(%s)-1)/4" % block_range(prefix, count, row))
        layer_range = "%s%d:%s%d" % (letters["DM_C"], row, letters["DM_L4"], row)
        formulas = {
            "ADMI_base": "=AVERAGE(%s)" % layer_range,
            "DM_min": "=MIN(%s)" % layer_range,
            "Penalty": "=1-alpha*MAX(0,delta-%s%d)" % (letters["DM_min"], row),
            "ADMI": "=%s%d*%s%d" % (letters["ADMI_base"], row, letters["Penalty"], row),
            "Level base": level_formula("%s%d" % (letters["ADMI_base"], row)),
            "Level ADMI": level_formula("%s%d" % (letters["ADMI"], row)),
            "Penalised": "=IF(%s%d<1,1,0)" % (letters["Penalty"], row),
            "Level changed": "=IF(%s%d<>%s%d,1,0)" % (letters["Level base"], row, letters["Level ADMI"], row),
            "Illusion of maturity": "=IF(AND(%s%d>=3,%s%d<=0.2),1,0)" % (letters["Level ADMI"], row, letters["DM_I"], row),
            "Scalar incl. DM_I": "=AVERAGE(%s%d:%s%d)" % (letters["DM_C"], row, letters["DM_I"], row),
            "Scalar level": level_formula("%s%d" % (letters["Scalar incl. DM_I"], row)),
            # #N/A leaves a point out of an Excel scatter chart
            "Figure 2: ADMI_base if level changed": "=IF(%s%d=1,%s%d,NA())" % (letters["Level changed"], row, letters["ADMI_base"], row),
            "Figure 2: ADMI_base if level kept": "=IF(%s%d=0,%s%d,NA())" % (letters["Level changed"], row, letters["ADMI_base"], row),
        }
        for header, formula in formulas.items():
            sheet.cell(row=row, column=headers.index(header) + 1, value=formula)

    def column_range(header):
        return "%s%d:%s%d" % (letters[header], first_row, letters[header], last_row)

    full_run = json.loads(SIMULATION_RESULTS.read_text(encoding="utf-8"))["scenarios"]["uniform"]
    summary = [
        (tr("Share of profiles with a penalty"), "=AVERAGE(%s)" % column_range("Penalised"), full_run["P1_calibrated"]["share_penalised"]),
        (tr("Share of profiles whose level changed (P1)"), "=AVERAGE(%s)" % column_range("Level changed"), full_run["P1_calibrated"]["share_level_changed"]),
        (tr("Mean reduction of the index"), "=AVERAGE(%s)-AVERAGE(%s)" % (column_range("ADMI_base"), column_range("ADMI")), full_run["P1_calibrated"]["mean_index_reduction"]),
        (tr("Share 'illusion of maturity' (P2)"), "=AVERAGE(%s)" % column_range("Illusion of maturity"), full_run["P2_integration_masking"]["share_illusion_of_maturity"]),
    ]
    sheet["A2"] = tr("Summary for this sample")
    sheet["A2"].font = BOLD
    sheet["E2"] = tr("Full run, 100 000 profiles (scripts/simulation_experiment.py)")
    sheet["E2"].font = BOLD
    for offset, (label, formula, full_value) in enumerate(summary):
        sheet.cell(row=3 + offset, column=1, value=label)
        sheet.cell(row=3 + offset, column=4, value=formula)
        sheet.cell(row=3 + offset, column=5, value=full_value)
    sheet.freeze_panes = sheet.cell(row=first_row, column=2)


SAMPLE_TRANSLATED = {"Level base", "Level ADMI", "Penalised", "Level changed", "Illusion of maturity", "Scalar incl. DM_I", "Scalar level",
                     "Figure 2: ADMI_base if level changed", "Figure 2: ADMI_base if level kept", "Penalty"}


def sample_label(header):
    """Header of an Experiment_sample column as written in the sheet; answer and DM_ columns keep their codes."""
    return tr(header) if header in SAMPLE_TRANSLATED else header


def sample_column(workbook, header):
    """Absolute reference to the data rows of one Experiment_sample column, plus its row bounds."""
    sheet = workbook[sheet_name("Experiment_sample")]
    header_row = 8
    for column in range(1, sheet.max_column + 1):
        if sheet.cell(row=header_row, column=column).value == sample_label(header):
            return column, header_row + 1, header_row + SAMPLE_SIZE
    raise KeyError(header)


def scatter_series(sheet_or_workbook, name, x_column, y_column, first_row, last_row, title, color, size=3, symbol="circle", line=False, dash=None):
    from openpyxl.chart import Reference, Series
    from openpyxl.chart.marker import Marker
    worksheet = sheet_or_workbook[sheet_name(name)]
    x_values = Reference(worksheet, min_col=x_column, min_row=first_row, max_row=last_row)
    y_values = Reference(worksheet, min_col=y_column, min_row=first_row, max_row=last_row)
    series = Series(y_values, x_values, title=title)
    if line:
        series.marker = Marker(symbol="none")
        series.graphicalProperties.line.solidFill = color
        series.graphicalProperties.line.width = 19050
        if dash:
            series.graphicalProperties.line.dashStyle = dash
    else:
        series.marker = Marker(symbol=symbol, size=size)
        series.marker.graphicalProperties.solidFill = color
        series.marker.graphicalProperties.line.solidFill = color
        series.graphicalProperties.line.noFill = True
    series.smooth = False
    return series


def style_axes(chart, x_title, y_title, x_range, y_range, x_step, y_step):
    chart.x_axis.title = x_title
    chart.y_axis.title = y_title
    chart.x_axis.scaling.min, chart.x_axis.scaling.max = x_range
    chart.y_axis.scaling.min, chart.y_axis.scaling.max = y_range
    chart.x_axis.majorUnit = x_step
    chart.y_axis.majorUnit = y_step
    chart.x_axis.number_format = chart.y_axis.number_format = "0.0"
    chart.x_axis.delete = chart.y_axis.delete = False
    chart.x_axis.majorGridlines = None
    chart.legend.position = "b"


def build_figure_band(workbook):
    """Figure 2: base index versus the weakest layer; the level drops across a bound inside the band of equation (7)."""
    from openpyxl.chart import ScatterChart
    sheet = workbook.create_sheet(sheet_name("Figure_2_band"))
    sheet["A1"] = tr("Figure 2. Where the weak-link penalty changes the maturity level (equation 7; points: first 1 000 profiles of the uniform scenario)")
    sheet["A1"].font = TITLE
    sheet["A2"] = tr("Between the solid line (bound b) and the dashed line (upper limit of equation 7) the penalty moves the profile one level down. Change alpha or delta on Parameters to see the band move.")
    headers = ["DM_min"] + [tr(template) % bound for bound in ("0.40", "0.60", "0.20") for template in ("Bound %s", "Upper limit for %s")]
    for column, header in enumerate(headers, start=1):
        sheet.cell(row=4, column=column, value=header).font = BOLD
    first_row = 5
    steps = 16  # DM_min from 0 to 0.375 in steps of 0.025
    for step in range(steps):
        row = first_row + step
        sheet.cell(row=row, column=1, value=round(step * 0.025, 3))
        for column, bound_index in ((2, 2), (4, 3), (6, 1)):
            sheet.cell(row=row, column=column, value="=INDEX(bounds,%d)" % bound_index)
            sheet.cell(row=row, column=column + 1,
                       value="=MIN(%s%d/(1-alpha*MAX(0,delta-A%d)),(4+A%d)/5)" % (sheet.cell(row=row, column=column).column_letter, row, row, row))
    last_row = first_row + steps - 1
    chart = ScatterChart()
    chart.title = tr("Penalty band and sample profiles")
    chart.style = 13
    chart.height, chart.width = 12, 20
    dm_min_column, sample_first, sample_last = sample_column(workbook, "DM_min")
    kept_column, _, _ = sample_column(workbook, "Figure 2: ADMI_base if level kept")
    changed_column, _, _ = sample_column(workbook, "Figure 2: ADMI_base if level changed")
    chart.series.append(scatter_series(workbook, "Experiment_sample", dm_min_column, kept_column, sample_first, sample_last, tr("level kept"), "BFBFBF", size=3))
    chart.series.append(scatter_series(workbook, "Experiment_sample", dm_min_column, changed_column, sample_first, sample_last, tr("level changed by the penalty"), "C0504D", size=4))
    for column, color, name in ((6, "7F7F7F", "0.20"), (2, "1F4E79", "0.40"), (4, "2E75B6", "0.60")):
        chart.series.append(scatter_series(workbook, "Figure_2_band", 1, column, first_row, last_row, tr("bound %s") % name, color, line=True))
        chart.series.append(scatter_series(workbook, "Figure_2_band", 1, column + 1, first_row, last_row, tr("upper limit %s") % name, color, line=True, dash="dash"))
    style_axes(chart, tr("DM_min (weakest layer)"), "ADMI_base", (0, 0.6), (0, 0.9), 0.1, 0.1)
    sheet.add_chart(chart, "I4")
    set_widths(sheet, {"A": 9, "B": 11, "C": 19, "D": 11, "E": 19, "F": 11, "G": 19})


def build_figure_map(workbook):
    """Figure 3: diagnostic map in the plane (ADMI, DM_I)."""
    from openpyxl.chart import ScatterChart
    sheet = workbook.create_sheet(sheet_name("Figure_3_map"))
    sheet["A1"] = tr("Figure 3. Diagnostic map (ADMI, DM_I): sample profiles and the situations of Table 4")
    sheet["A1"].font = TITLE
    sheet["A2"] = tr("Lower right zone (ADMI above the level-3 bound, DM_I at most 0.20) is the 'illusion of maturity': equipped layers without integration.")
    sheet["A4"], sheet["B4"], sheet["C4"] = tr("Situation"), "ADMI", "DM_I"
    for cell in sheet[4]:
        cell.font = BOLD
    situations = [tr(name) for name in ["Compliance trap", "External EU pressure", "Absorptive limiter", "Aware mismatch", "Illusion of maturity"]]
    for offset, name in enumerate(situations):
        row = 5 + offset
        sheet.cell(row=row, column=1, value=name)
        sheet.cell(row=row, column=2, value="=%sL%d" % (ref("Diagnostic_situations"), 4 + offset))
        sheet.cell(row=row, column=3, value="=%sG%d" % (ref("Diagnostic_situations"), 4 + offset))
    sheet["A12"], sheet["B12"], sheet["C12"] = tr("Guide lines"), "x", "y"
    for cell in sheet[12]:
        cell.font = BOLD
    guide_rows = [(tr("level-3 bound"), "=INDEX(bounds,2)", 0), (tr("level-3 bound"), "=INDEX(bounds,2)", 1),
                  ("DM_I = 0.20", 0, 0.2), ("DM_I = 0.20", 1, 0.2)]
    for offset, (name, x_value, y_value) in enumerate(guide_rows):
        sheet.cell(row=13 + offset, column=1, value=name)
        sheet.cell(row=13 + offset, column=2, value=x_value)
        sheet.cell(row=13 + offset, column=3, value=y_value)
    chart = ScatterChart()
    chart.title = tr("Diagnostic map (ADMI, DM_I)")
    chart.style = 13
    chart.height, chart.width = 12, 20
    admi_column, sample_first, sample_last = sample_column(workbook, "ADMI")
    integration_column, _, _ = sample_column(workbook, "DM_I")
    chart.series.append(scatter_series(workbook, "Experiment_sample", admi_column, integration_column, sample_first, sample_last, tr("sample profiles"), "D9D9D9", size=3))
    colors = ["C0504D", "F79646", "9BBB59", "8064A2", "1F4E79"]
    for offset, (name, color) in enumerate(zip(situations, colors)):
        chart.series.append(scatter_series(workbook, "Figure_3_map", 2, 3, 5 + offset, 5 + offset, name, color, size=9, symbol="diamond"))
    chart.series.append(scatter_series(workbook, "Figure_3_map", 2, 3, 13, 14, tr("level-3 bound"), "595959", line=True, dash="dash"))
    chart.series.append(scatter_series(workbook, "Figure_3_map", 2, 3, 15, 16, "DM_I = 0.20", "595959", line=True, dash="dash"))
    style_axes(chart, "ADMI", tr("DM_I (integration)"), (0, 1), (0, 1), 0.2, 0.2)
    sheet.add_chart(chart, "E4")
    set_widths(sheet, {"A": 22, "B": 9, "C": 9})


def build(code):
    set_language(code)
    output_path, readme_path = WORKBOOKS[code]
    workbook = Workbook()
    build_readme(workbook.active)
    workbook.active.title = sheet_name("README")
    build_parameters(workbook)
    build_calculator(workbook, "Calculator", WORKED_EXAMPLE_ANSWERS, "ADMI calculator: enter answers 1-5 in the yellow cells")
    build_calculator(workbook, "Worked_example", WORKED_EXAMPLE_ANSWERS, "Worked example of Section 4 (ADMI_base 0.433, level 3 -> ADMI 0.383, level 2)")
    build_data_sheets(workbook)
    build_statistics_mode(workbook)
    build_topsis(workbook)
    workbook.move_sheet(sheet_name("Statistics_mode"), offset=-2)
    workbook.move_sheet(sheet_name("TOPSIS"), offset=-2)
    build_diagnostic_situations(workbook)
    build_boundary_band(workbook)
    build_experiment_sample(workbook)
    build_figure_band(workbook)
    build_figure_map(workbook)
    for row in workbook[sheet_name("README")].iter_rows():
        for cell in row:
            cell.alignment = WRAP
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(str(output_path))
    write_readme_markdown(readme_path)
    print("Saved:", output_path)
    print("Saved:", readme_path)


def main():
    for code in WORKBOOKS:
        build(code)


if __name__ == "__main__":
    main()
